"""Stage 1 - splits a PDF/DOCX/TXT into clauses. OCR fallback for scanned pages."""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from pathlib import Path

from backend.app.core.config import Settings, get_settings
from backend.app.core.schemas import BoundingBox, ParsedClause
from backend.app.services.text_utils import normalise_text

logger = logging.getLogger(__name__)
_NUMBERED = re.compile(
    r"^\s*(?P<num>(?:"
    r"\((?:\d+(?:\.\d+)*|[A-Z]|[IVXLC]+)\)"
    r"|\d+(?:\.\d+)+(?:[.)])?"
    r"|\d+[.)]"
    r"|[A-Z][.)]"
    r"|[IVXLC]+[.)]"
    r"))\s*(?P<body>.*)$",
    re.I,
)
_HEADING = re.compile(r"^[A-Z][A-Z0-9 /&(),.'-]{3,100}$")


@dataclass
class _Block:
    text: str
    page_number: int | None
    block_number: int | None
    method: str
    bbox: tuple[float, float, float, float] | None = None
    page_height: float | None = None


class Stage1DocumentParser:
    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or get_settings()

    def parse(self, file_path: Path) -> list[ParsedClause]:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            blocks = self._parse_pdf(file_path)
        elif suffix == ".docx":
            blocks = self._parse_docx(file_path)
        elif suffix == ".txt":
            blocks = [_Block(file_path.read_text(encoding="utf-8", errors="replace"), None, 0, "txt")]
        else:
            raise ValueError(f"Unsupported document type: {suffix}")
        blocks = self._remove_repeated_margin_blocks(blocks)
        clauses = self._segment(blocks)
        if len(clauses) > self.settings.parser_max_clauses:
            raise ValueError(f"Document produced {len(clauses)} clauses; maximum is {self.settings.parser_max_clauses}")
        if not clauses:
            raise ValueError("No clauses could be extracted")
        return clauses

    def _parse_pdf(self, file_path: Path) -> list[_Block]:
        import fitz

        output = []
        with fitz.open(file_path) as document:
            if document.needs_pass:
                raise ValueError("Encrypted PDFs are not supported")
            if document.page_count > self.settings.parser_max_pages:
                raise ValueError(f"PDF has {document.page_count} pages; maximum is {self.settings.parser_max_pages}")
            for page_index, page in enumerate(document):
                page_number = page_index + 1
                page_blocks, visible = [], 0
                for block_index, block in enumerate(page.get_text("blocks")):
                    text = str(block[4]).strip()
                    if not text:
                        continue
                    visible += len(re.sub(r"\s+", "", text))
                    page_blocks.append(
                        _Block(
                            text,
                            page_number,
                            block_index,
                            "pymupdf",
                            (float(block[0]), float(block[1]), float(block[2]), float(block[3])),
                            float(page.rect.height),
                        )
                    )
                if visible < self.settings.parser_ocr_min_chars_per_page:
                    ocr = self._ocr_page(page)
                    if ocr.strip():
                        output.append(_Block(ocr, page_number, 0, "ocr"))
                        continue
                    fallback = self._pdfplumber_page(file_path, page_index)
                    if fallback.strip():
                        output.append(_Block(fallback, page_number, 0, "pdfplumber"))
                        continue
                output.extend(page_blocks)
        return output

    def _ocr_page(self, page) -> str:
        try:
            import pytesseract
            from PIL import Image

            pix = page.get_pixmap(dpi=self.settings.parser_ocr_dpi, alpha=False)
            return pytesseract.image_to_string(Image.open(io.BytesIO(pix.tobytes("png"))), lang="eng")
        except Exception as exc:
            logger.warning("OCR failed on page %s: %s", page.number + 1, exc)
            return ""

    @staticmethod
    def _pdfplumber_page(file_path: Path, page_index: int) -> str:
        try:
            import pdfplumber

            with pdfplumber.open(file_path) as pdf:
                return pdf.pages[page_index].extract_text(layout=True) or ""
        except Exception as exc:
            logger.warning("pdfplumber failed on page %s: %s", page_index + 1, exc)
            return ""

    @staticmethod
    def _parse_docx(file_path: Path) -> list[_Block]:
        from docx import Document

        document = Document(file_path)
        output, index = [], 0
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                output.append(_Block(paragraph.text.strip(), None, index, "docx"))
                index += 1
        for table in document.tables:
            for row in table.rows:
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    output.append(_Block(text, None, index, "docx"))
                    index += 1
        return output

    @staticmethod
    def _remove_repeated_margin_blocks(blocks: list[_Block]) -> list[_Block]:
        """Remove repeated PDF headers/footers while retaining body clauses."""
        from collections import defaultdict
        import math

        pages = {block.page_number for block in blocks if block.page_number is not None}
        if len(pages) < 2:
            return blocks
        occurrences: dict[str, set[int]] = defaultdict(set)
        for block in blocks:
            if block.page_number is None or block.bbox is None or block.page_height is None:
                continue
            top = block.bbox[1] <= 72
            bottom = block.bbox[3] >= block.page_height - 72
            clean = normalise_text(block.text)
            if (top or bottom) and clean and len(clean) <= 160:
                occurrences[clean].add(block.page_number)
        threshold = max(2, math.ceil(len(pages) * 0.5))
        repeated = {text for text, seen_pages in occurrences.items() if len(seen_pages) >= threshold}
        if not repeated:
            return blocks
        return [block for block in blocks if normalise_text(block.text) not in repeated]

    def _segment(self, blocks: list[_Block]) -> list[ParsedClause]:
        candidates, heading = [], None
        for block in blocks:
            lines = [line.strip() for line in re.split(r"\r?\n+", block.text) if line.strip()]
            buffer = ""
            for line in lines:
                if self._is_heading(line):
                    if buffer:
                        candidates.append((buffer, block, heading))
                        buffer = ""
                    heading = normalise_text(line)
                    continue
                numbered = _NUMBERED.match(line)
                if numbered:
                    if buffer:
                        candidates.append((buffer, block, heading))
                    buffer = f"{numbered.group('num')} {numbered.group('body').strip()}".strip()
                else:
                    buffer = f"{buffer} {line}".strip()
            if buffer:
                candidates.append((buffer, block, heading))
        merged, pending = [], None
        for text, block, heading in candidates:
            clean = normalise_text(text)
            if not clean:
                continue
            if (
                re.fullmatch(r"(?:\d+(?:\.\d+)*|[A-Z]|[IVXLC]+)[.)]", clean, re.I)
                or len(clean.split()) < self.settings.parser_min_clause_words
            ):
                pending = (clean, block, heading)
                continue
            if pending:
                clean = f"{pending[0]} {clean}"
                block = pending[1]
                heading = heading or pending[2]
                pending = None
            merged.append((clean, block, heading))
        if pending and merged:
            text, block, heading = merged[-1]
            merged[-1] = (f"{text} {pending[0]}", block, heading)
        results = []
        for index, (text, block, heading) in enumerate(merged, 1):
            bbox = (
                BoundingBox(x0=block.bbox[0], y0=block.bbox[1], x1=block.bbox[2], y1=block.bbox[3])
                if block.bbox
                else None
            )
            results.append(
                ParsedClause(
                    clause_id=f"C{index:04d}",
                    text=text,
                    page_number=block.page_number,
                    block_number=block.block_number,
                    section_heading=heading,
                    extraction_method=block.method,
                    bounding_box=bbox,
                )
            )
        return results

    @staticmethod
    def _is_heading(line: str) -> bool:
        clean = normalise_text(line)
        if len(clean.split()) > 12:
            return False
        return bool(_HEADING.fullmatch(clean)) or clean.endswith(":")
