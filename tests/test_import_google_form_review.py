"""google forms gives one row per reviewer, but scoring needs one row per clause - these check the reshape doesn't lose or duplicate anything"""

from __future__ import annotations

import csv

import openpyxl
import pytest

from scripts.import_google_form_review import (
    fill_review_workbook,
    parse_google_form_responses,
    write_scoring_csv,
)

RECORD_A = "cuad::example_a.pdf::1::10::20::abc123::supported"
RECORD_B = "cuad::example_b.pdf::2::30::40::def456::unsupported"


def _build_form_workbook(path):
    workbook = openpyxl.Workbook()
    responses = workbook.active
    responses.title = "Form responses 1"
    responses.append(
        [
            "Timestamp",
            "Email address",
            "Reviewer Name / ID (e.g., Reviewer 1, Reviewer 2)",
            f"[{RECORD_A}] Risk Category (pick one):",
            f"[{RECORD_A}] Does the Plain-English Version Mean the Same Thing?",
            f"[{RECORD_A}] Is Anything Important Missing, Wrong, or Changed?",
            f"[{RECORD_A}] If Yes, what's missing or wrong? (leave blank if No)",
            f"[{RECORD_A}] Comments (optional):",
            f"[{RECORD_B}] Risk Category (pick one):",
            f"[{RECORD_B}] Does the Plain-English Version Mean the Same Thing?",
            f"[{RECORD_B}] Is Anything Important Missing, Wrong, or Changed?",
            f"[{RECORD_B}] If Yes, what's missing or wrong? (leave blank if No)",
            f"[{RECORD_B}] Comments (optional):",
        ]
    )
    responses.append(
        ["2026-07-30", "r1@example.com", "Reviewer 1", "Safe", "Yes", "No", "", "Looks fine",
         "Hidden Fee", "No", "Yes", "Missed a fee amount", ""]
    )
    responses.append(
        ["2026-07-30", "r2@example.com", "Reviewer 2", "Safe", "Yes", "No", "", "",
         "Hidden Fee", "No", "No", "", "Agrees with reviewer 1"]
    )

    review = workbook.create_sheet("Review")
    review.append(
        [
            "ID", "Original Contract Clause", "Plain-English Version", "Supporting Regulation Text (if any)",
            "Reviewer 1: Risk Category (pick one)",
            "Reviewer 1: Does the Plain-English Version Mean the Same Thing? (Yes / No)",
            "Reviewer 1: Is Anything Important Missing, Wrong, or Changed? (Yes / No)",
            "Reviewer 1: If Yes, what's missing or wrong? (leave blank if No)",
            "Reviewer 1: Comments (optional)",
            "Reviewer 2: Risk Category (pick one)",
            "Reviewer 2: Does the Plain-English Version Mean the Same Thing? (Yes / No)",
            "Reviewer 2: Is Anything Important Missing, Wrong, or Changed? (Yes / No)",
            "Reviewer 2: If Yes, what's missing or wrong? (leave blank if No)",
            "Reviewer 2: Comments (optional)",
        ]
    )
    review.append([RECORD_A, "clause a", "simplified a", "reg a"])
    review.append([RECORD_B, "clause b", "simplified b", "reg b"])
    workbook.save(path)


@pytest.fixture
def form_workbook(tmp_path):
    path = tmp_path / "review.xlsx"
    _build_form_workbook(path)
    return path


def test_parse_extracts_both_reviewers_for_both_clauses(form_workbook):
    answers, diagnostics = parse_google_form_responses(form_workbook)
    assert set(answers.keys()) == {RECORD_A, RECORD_B}
    assert answers[RECORD_A]["reviewer_1_risk_category"] == "Safe"
    assert answers[RECORD_A]["reviewer_2_missing_or_changed"] == "No"
    assert answers[RECORD_B]["reviewer_1_missing_or_changed_details"] == "Missed a fee amount"
    assert diagnostics["reviewer_labels_seen"] == ["Reviewer 1", "Reviewer 2"]
    assert diagnostics["unrecognised_columns"] == []


def test_write_scoring_csv_preserves_reference_order_and_flags_gaps(form_workbook, tmp_path):
    answers, _ = parse_google_form_responses(form_workbook)
    out_csv = tmp_path / "scoring.csv"
    result = write_scoring_csv(answers, [RECORD_A, RECORD_B, "unreviewed::clause"], out_csv, n_reviewers=2)
    assert result["missing_record_ids"] == ["unreviewed::clause"]
    with out_csv.open(encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    assert [row["record_id"] for row in rows] == [RECORD_A, RECORD_B, "unreviewed::clause"]
    assert rows[0]["reviewer_2_risk_category"] == "Safe"
    assert rows[1]["reviewer_2_risk_category"] == "Hidden Fee"
    assert rows[2]["reviewer_1_risk_category"] == ""


def test_fill_review_workbook_writes_new_file_without_touching_original(form_workbook, tmp_path):
    answers, _ = parse_google_form_responses(form_workbook)
    out_path = tmp_path / "filled.xlsx"
    result = fill_review_workbook(form_workbook, answers, out_path)
    assert result["filled_cells"] == 20
    assert result["unmatched_rows_in_review_sheet"] == 0

    original_review = openpyxl.load_workbook(form_workbook)["Review"]
    assert original_review.cell(row=2, column=5).value is None

    filled_review = openpyxl.load_workbook(out_path)["Review"]
    assert filled_review.cell(row=2, column=5).value == "Safe"
    assert filled_review.cell(row=3, column=10).value == "Hidden Fee"
