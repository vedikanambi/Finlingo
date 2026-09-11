"""Reshapes Google Forms review responses (one row per reviewer) into the
row-per-clause format score_review_file expects. Never touches the original
spreadsheet - always writes to a new output path."""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl

_QUESTION_FIELD = {
    "Risk Category": "risk_category",
    "Mean the Same Thing": "meaning_preserved",
    "Missing, Wrong, or Changed": "missing_or_changed",
    "what's missing or wrong": "missing_or_changed_details",
    "Comments": "comments",
}
_HEADER_RE = re.compile(r"^\[(?P<record_id>.+)\]\s*(?P<question>.+)$")
_REVIEWER_NUMBER_RE = re.compile(r"(\d+)")


def _field_for_question(question: str) -> str | None:
    for needle, field in _QUESTION_FIELD.items():
        if needle in question:
            return field
    return None


def _find_response_sheet(workbook: openpyxl.Workbook) -> str:
    """pick the sheet with data - google forms leaves stale empty ones behind if you rebuild the form"""
    candidates = [name for name in workbook.sheetnames if name.startswith("Form responses")]
    if not candidates:
        raise ValueError(f"No 'Form responses *' sheet found. Sheets present: {workbook.sheetnames}")
    best = max(candidates, key=lambda name: workbook[name].max_row)
    if workbook[best].max_row < 2:
        raise ValueError(f"'{best}' (and all other Form responses sheets) has no response rows yet.")
    return best


def _reviewer_column_index(header_row: tuple, sheet_name: str) -> int:
    for i, value in enumerate(header_row):
        if value and "Reviewer Name" in str(value):
            return i
    raise ValueError(f"Could not find the 'Reviewer Name / ID' column in '{sheet_name}'.")


def parse_google_form_responses(xlsx_path: Path) -> tuple[dict[str, dict[str, str]], dict[str, object]]:
    """returns (per_record_id reviewer answers, diagnostics)"""
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet_name = _find_response_sheet(workbook)
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]
    reviewer_col = _reviewer_column_index(header, sheet_name)

    column_map: dict[int, tuple[str, str]] = {}
    unrecognised_columns: list[str] = []
    for i, value in enumerate(header):
        if value is None or i == reviewer_col:
            continue
        match = _HEADER_RE.match(str(value))
        if not match:
            continue
        field = _field_for_question(match.group("question"))
        if field is None:
            unrecognised_columns.append(str(value))
            continue
        column_map[i] = (match.group("record_id").strip(), field)

    answers: dict[str, dict[str, str]] = {}
    reviewer_labels_seen: list[str] = []
    for data_row in rows[1:]:
        if all(v is None for v in data_row):
            continue
        reviewer_label = str(data_row[reviewer_col] or "").strip()
        number_match = _REVIEWER_NUMBER_RE.search(reviewer_label)
        if not number_match:
            raise ValueError(
                f"Could not parse a reviewer number out of '{reviewer_label}' -- expected e.g. 'Reviewer 2'."
            )
        reviewer_number = int(number_match.group(1))
        reviewer_labels_seen.append(reviewer_label)
        for col_index, (record_id, field) in column_map.items():
            value = data_row[col_index]
            answers.setdefault(record_id, {})[f"reviewer_{reviewer_number}_{field}"] = (
                "" if value is None else str(value).strip()
            )

    diagnostics = {
        "source_sheet": sheet_name,
        "reviewer_labels_seen": reviewer_labels_seen,
        "unrecognised_columns": unrecognised_columns,
        "record_ids_found": sorted(answers.keys()),
    }
    return answers, diagnostics


def write_scoring_csv(
    answers: dict[str, dict[str, str]],
    reference_record_ids: list[str],
    out_path: Path,
    n_reviewers: int,
) -> dict:
    fieldnames = ["record_id"]
    fields = ["risk_category", "meaning_preserved", "missing_or_changed", "missing_or_changed_details", "comments"]
    for i in range(1, n_reviewers + 1):
        fieldnames.extend(f"reviewer_{i}_{f}" for f in fields)

    missing_record_ids = [rid for rid in reference_record_ids if rid not in answers]
    extra_record_ids = [rid for rid in answers if rid not in reference_record_ids]

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for record_id in reference_record_ids:
            row = {"record_id": record_id, **answers.get(record_id, {})}
            writer.writerow(row)

    return {
        "rows_written": len(reference_record_ids),
        "missing_record_ids": missing_record_ids,
        "extra_record_ids_ignored": extra_record_ids,
    }


def fill_review_workbook(
    xlsx_path: Path,
    answers: dict[str, dict[str, str]],
    out_path: Path,
) -> dict:
    """fills reviewer answers into a copy of the workbook, never the original"""
    workbook = openpyxl.load_workbook(xlsx_path)
    if "Review" not in workbook.sheetnames:
        raise ValueError("Expected a 'Review' sheet in the workbook.")
    sheet = workbook["Review"]
    header = [cell.value for cell in sheet[1]]

    field_by_header_fragment = {
        "Risk Category": "risk_category",
        "Mean the Same Thing": "meaning_preserved",
        "Missing, Wrong, or Changed": "missing_or_changed",
        "what's missing or wrong": "missing_or_changed_details",
        "Comments": "comments",
    }
    reviewer_col_re = re.compile(r"^Reviewer (\d+):")
    column_targets: dict[int, tuple[int, str]] = {}
    for col_index, value in enumerate(header, start=1):
        if not value:
            continue
        rmatch = reviewer_col_re.match(str(value))
        if not rmatch:
            continue
        field = _field_for_question(str(value))
        if field is None:
            continue
        column_targets[col_index] = (int(rmatch.group(1)), field)

    id_col = 1
    filled_cells = 0
    unmatched_rows = 0
    for row_index in range(2, sheet.max_row + 1):
        record_id = str(sheet.cell(row=row_index, column=id_col).value or "").strip()
        if not record_id:
            continue
        record_answers = answers.get(record_id)
        if not record_answers:
            unmatched_rows += 1
            continue
        for col_index, (reviewer_number, field) in column_targets.items():
            key = f"reviewer_{reviewer_number}_{field}"
            if key in record_answers:
                sheet.cell(row=row_index, column=col_index, value=record_answers[key])
                filled_cells += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)
    return {"filled_cells": filled_cells, "unmatched_rows_in_review_sheet": unmatched_rows}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path, help="Path to the workbook with the Google Forms response sheet(s)")
    parser.add_argument(
        "--scoring-csv-out",
        type=Path,
        default=Path("reports/human_review_completed_v1.csv"),
        help="CSV consumed by `python main.py score-review` (default: reports/human_review_completed_v1.csv)",
    )
    parser.add_argument(
        "--filled-xlsx-out",
        type=Path,
        default=None,
        help="If set, writes a NEW xlsx with reviewer answers filled into the 'Review' sheet (original is never modified)",
    )
    parser.add_argument("--n-reviewers", type=int, default=3)
    args = parser.parse_args()

    answers, diagnostics = parse_google_form_responses(args.xlsx)
    print(f"Read {len(answers)} clauses' worth of answers from sheet '{diagnostics['source_sheet']}'")
    print(f"Reviewer labels seen: {diagnostics['reviewer_labels_seen']}")
    if diagnostics["unrecognised_columns"]:
        print(f"WARNING: {len(diagnostics['unrecognised_columns'])} column(s) did not match a known question type:")
        for col in diagnostics["unrecognised_columns"]:
            print(f"  - {col}")

    workbook = openpyxl.load_workbook(args.xlsx, data_only=True)
    reference_record_ids = [
        str(row[0].value).strip()
        for row in workbook["Review"].iter_rows(min_row=2)
        if row[0].value
    ]

    csv_result = write_scoring_csv(answers, reference_record_ids, args.scoring_csv_out, args.n_reviewers)
    print(f"\nWrote scoring CSV to {args.scoring_csv_out}: {csv_result}")
    if csv_result["missing_record_ids"]:
        print(f"WARNING: {len(csv_result['missing_record_ids'])} reference clause(s) have no reviewer answers yet.")

    if args.filled_xlsx_out:
        fill_result = fill_review_workbook(args.xlsx, answers, args.filled_xlsx_out)
        print(f"\nWrote filled workbook to {args.filled_xlsx_out}: {fill_result}")


if __name__ == "__main__":
    main()
